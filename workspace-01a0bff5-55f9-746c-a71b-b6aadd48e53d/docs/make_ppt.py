"""Generate the 10-slide hackathon pitch deck for HireMind AI (PS03)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG = RGBColor(0x0B, 0x10, 0x20)
CARD = RGBColor(0x15, 0x1C, 0x33)
LINE = RGBColor(0x26, 0x30, 0x52)
WHITE = RGBColor(0xE7, 0xEC, 0xF7)
MUTED = RGBColor(0x8B, 0x96, 0xB5)
INDIGO = RGBColor(0x81, 0x8C, 0xF8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
EMERALD = RGBColor(0x34, 0xD3, 0x99)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
RED = RGBColor(0xF8, 0x71, 0x71)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def box(s, l, t, w, h, fill=CARD, border=LINE, radius=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def tb(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tx = s.shapes.add_textbox(l, t, w, h)
    tf = tx.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def run(p, text, size=18, color=WHITE, bold=False, italic=False, font="Segoe UI"):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    return r


def title_bar(s, num, text, accent=INDIGO):
    p = tb(s, Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.9)).paragraphs[0]
    run(p, num + "  ", size=30, color=accent, bold=True)
    run(p, text, size=30, color=WHITE, bold=True)
    # underline accent
    bar = s.shapes.add_shape(1, Inches(0.6), Inches(1.18), Inches(2.2), Pt(4))  # rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False


def footer(s, n):
    p = tb(s, Inches(0.55), Inches(7.05), Inches(9), Inches(0.4)).paragraphs[0]
    run(p, "HireMind AI · PS03 Multi-Agent Resume Screening & Job Matching", size=10, color=MUTED)
    p2 = tb(s, Inches(12.3), Inches(7.05), Inches(0.6), Inches(0.4)).paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    run(p2, str(n), size=10, color=MUTED)


def bullets(s, items, l=0.7, t=1.6, w=11.9, h=5.2, size=18, gap=10, color=WHITE, bullet_color=CYAN):
    tf = tb(s, Inches(l), Inches(t), Inches(w), Inches(h))
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        run(p, "▸  ", size=size, color=bullet_color, bold=True)
        if isinstance(it, tuple):
            run(p, it[0], size=size, color=color, bold=True)
            run(p, "  " + it[1], size=size - 1, color=MUTED)
        else:
            run(p, it, size=size, color=color)
    return tf


def simple_table(s, rows, cols, l, t, w, h, header=True):
    gf = s.shapes.add_table(rows, cols, l, t, w, h)
    tbl = gf.table
    tbl.first_row = header
    return tbl


def cell_set(tbl, r, c, text, size=14, color=WHITE, bold=False, fill=None, align=PP_ALIGN.LEFT):
    cell = tbl.cell(r, c)
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run(p, text, size=size, color=color, bold=bold)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


# ============================================================
# SLIDE 1 — TITLE
# ============================================================
s = slide()
logo = box(s, Inches(5.9), Inches(0.9), Inches(1.5), Inches(1.5), fill=RGBColor(0x4F, 0x46, 0xE5), border=RGBColor(0x4F, 0x46, 0xE5))
p = logo.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "🧠", size=44)
logo.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tb(s, Inches(0.5), Inches(2.65), Inches(12.3), Inches(1.2)).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "HireMind AI", size=54, color=WHITE, bold=True)

p = tb(s, Inches(1), Inches(3.9), Inches(11.3), Inches(0.7)).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "Multi-Agent Resume Screening & Job Matching Platform", size=24, color=INDIGO, bold=True)

p = tb(s, Inches(1), Inches(4.75), Inches(11.3), Inches(0.6)).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "Problem Statement PS03  ·  NLP + LLM + Agentic AI", size=16, color=MUTED)

box(s, Inches(3.4), Inches(5.6), Inches(6.5), Inches(0.75), fill=RGBColor(0x05, 0x2E, 0x22), border=RGBColor(0x06, 0x5F, 0x46))
p = tb(s, Inches(3.4), Inches(5.72), Inches(6.5), Inches(0.6), anchor=MSO_ANCHOR.MIDDLE).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "LLM extracts  ·  Deterministic engine decides  ·  Fully explainable", size=15, color=EMERALD, bold=True)

p = tb(s, Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.5)).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "[YOUR TEAM NAME]  ·  [HACKATHON NAME]  ·  11 Sept 2026", size=13, color=MUTED)

# ============================================================
# SLIDE 2 — PROBLEM
# ============================================================
s = slide()
title_bar(s, "01", "The Problem: Hiring Screening Doesn't Scale")
bullets(s, [
    ("7 seconds —", "average first-pass time a recruiter spends on a resume; hundreds of applications per opening"),
    ("Inconsistent screens —", "different screeners, different outcomes; strong candidates get missed"),
    ("Vocabulary mismatch —", "'React.js' vs 'React' vs 'ReactJS' breaks keyword ATS systems"),
    ("Opaque black boxes —", "pure-LLM or embedding scores can't justify a ranking to a candidate or auditor"),
    ("No feedback loop —", "bare scores don't tell recruiters what's missing or how close a candidate is"),
])
box(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.85), fill=RGBColor(0x17, 0x20, 0x3F), border=INDIGO)
p = tb(s, Inches(0.95), Inches(6.14), Inches(11.4), Inches(0.6), anchor=MSO_ANCHOR.MIDDLE).paragraphs[0]
run(p, "We need automated matching that is fast, consistent, AND explainable.", size=17, color=INDIGO, bold=True)
footer(s, 2)

# ============================================================
# SLIDE 3 — SOLUTION
# ============================================================
s = slide()
title_bar(s, "02", "Our Solution: Five Cooperating AI Agents")
agents = [
    ("① Resume Agent", "parses resumes (txt/pdf/docx) into structured candidate profiles", CYAN),
    ("② Job Agent", "extracts must-have vs nice-to-have skills, experience & education", CYAN),
    ("③ Matching Agent", "deterministic weighted 0–100 fit score — explainable & reproducible", EMERALD),
    ("④ Skill Gap Agent", "critical vs stretch gaps + adjacent strengths + upskilling plan", AMBER),
    ("⑤ Recruiter Agent", "human-style summary with interview / hold / pass recommendation", AMBER),
]
tf = tb(s, Inches(0.7), Inches(1.6), Inches(11.9), Inches(4.3))
for i, (name, desc, col) in enumerate(agents):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(14)
    run(p, name, size=21, color=col, bold=True)
    run(p, "  —  " + desc, size=17, color=WHITE)
box(s, Inches(0.7), Inches(5.95), Inches(11.9), Inches(0.9), fill=RGBColor(0x05, 0x2E, 0x22), border=EMERALD)
p = tb(s, Inches(0.95), Inches(6.1), Inches(11.4), Inches(0.65), anchor=MSO_ANCHOR.MIDDLE).paragraphs[0]
run(p, "Core principle: \"LLM extracts, never decides\" — language models parse & write; a transparent engine ranks.", size=16, color=EMERALD, bold=True)
footer(s, 3)

# ============================================================
# SLIDE 4 — ARCHITECTURE
# ============================================================
s = slide()
title_bar(s, "03", "Architecture & Technical Flow")
s.shapes.add_picture("/home/user/docs/architecture.png", Inches(1.15), Inches(1.35), width=Inches(11.0))
p = tb(s, Inches(0.7), Inches(6.85), Inches(12), Inches(0.4)).paragraphs[0]
run(p, "Every agent's input/output is exposed in the UI via an Agent Trace panel — full auditability per decision.", size=13, color=MUTED, italic=True)
footer(s, 4)

# ============================================================
# SLIDE 5 — SCORING MODEL
# ============================================================
s = slide()
title_bar(s, "04", "Explainable Scoring Model (Matching Agent)")
tbl = simple_table(s, 6, 3, Inches(0.7), Inches(1.55), Inches(7.4), Inches(3.4))
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(1.1)
tbl.columns[2].width = Inches(3.7)
data = [
    ("Dimension", "Weight", "How it's computed"),
    ("Must-have skills", "45%", "fraction matched (fuzzy + taxonomy)"),
    ("Nice-to-have skills", "15%", "fraction matched"),
    ("Experience", "20%", "1.0 if ≥ required yrs, scaled below"),
    ("Education", "10%", "degree-rank comparison"),
    ("Keyword affinity", "10%", "JD keywords found in resume text"),
]
for r, row in enumerate(data):
    for c, val in enumerate(row):
        hdr = r == 0
        cell_set(tbl, r, c, val, size=14, bold=hdr,
                 color=WHITE if hdr else WHITE,
                 fill=RGBColor(0x17, 0x20, 0x3F) if hdr else CARD)
bullets(s, [
    ("Fuzzy matching —", "'React.js' = 'React' (difflib + 130-skill alias taxonomy)"),
    ("Verdicts —", "≥80 Strong · 60–79 Good · 40–59 Partial · <40 Weak"),
    ("Reproducible —", "same input ⇒ identical ranking, every run"),
    ("No demographics —", "only skills / experience / education enter scoring"),
], l=8.4, t=1.55, w=4.3, h=4.2, size=15, gap=14)
box(s, Inches(0.7), Inches(5.3), Inches(11.9), Inches(1.35), fill=RGBColor(0x17, 0x20, 0x3F), border=LINE)
tf = tb(s, Inches(0.95), Inches(5.42), Inches(11.4), Inches(1.15))
p = tf.paragraphs[0]
run(p, "Pipeline: ", size=15, color=CYAN, bold=True)
run(p, "Upload → Extraction (Gemini / offline NLP) → Weighted Scoring → Gap Analysis → Recruiter Summary → Ranked List + Agent Trace", size=15, color=WHITE)
p2 = tf.add_paragraph()
run(p2, "Dual mode: works with a Gemini API key OR fully offline — the demo never breaks on stage.", size=14, color=AMBER, bold=True)
footer(s, 5)

# ============================================================
# SLIDE 6 — TECH STACK
# ============================================================
s = slide()
title_bar(s, "05", "Technology Stack")
tbl = simple_table(s, 7, 3, Inches(0.7), Inches(1.55), Inches(11.9), Inches(4.4))
tbl.columns[0].width = Inches(2.4)
tbl.columns[1].width = Inches(3.9)
tbl.columns[2].width = Inches(5.6)
rows = [
    ("Layer", "Technology", "Why"),
    ("Frontend", "React 18 + Vite", "agent cards, ranked lists, explainable detail panels, hot-reload dev"),
    ("Backend", "Python 3 + FastAPI (Uvicorn)", "async REST API, auto OpenAPI docs, serves the React build"),
    ("LLM", "Google Gemini 2.5 Flash", "JSON-mode extraction & summarisation; user-supplied key, never hard-coded"),
    ("Offline NLP", "Regex + skill taxonomy + difflib", "deterministic fallback & alias normalisation ('k8s' → Kubernetes)"),
    ("Parsing", "pypdf · python-docx", "real-world resume formats"),
    ("Concurrency", "ThreadPoolExecutor", "parallel recruiter summaries for speed"),
]
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        hdr = r == 0
        cell_set(tbl, r, c, val, size=15 if hdr else 14, bold=hdr,
                 fill=RGBColor(0x17, 0x20, 0x3F) if hdr else CARD)
p = tb(s, Inches(0.7), Inches(6.15), Inches(11.9), Inches(0.6)).paragraphs[0]
run(p, "Repo: backend/ (agents) · frontend/ (React) · data/ (demo set) · docs/ (this deck + report) · test_pipeline.py", size=13, color=MUTED, italic=True)
footer(s, 6)

# ============================================================
# SLIDE 7 — RESULTS
# ============================================================
s = slide()
title_bar(s, "06", "Live Demo: 10 Resumes × 3 Jobs → Ranked Lists")
tbl = simple_table(s, 4, 4, Inches(0.7), Inches(1.55), Inches(11.9), Inches(2.6))
tbl.columns[0].width = Inches(3.1)
tbl.columns[1].width = Inches(3.0)
tbl.columns[2].width = Inches(3.0)
tbl.columns[3].width = Inches(2.8)
rows = [
    ("Job Description", "#1 Ranked", "#2", "#3"),
    ("Senior ML Engineer", "Aarav Sharma · 95 Strong", "Priya Nair · 69 Good", "Divya Reddy · 59 Partial"),
    ("Full-Stack Developer", "Sneha Kulkarni · 100 Strong", "Rohan Mehta · 81 Strong", "Vikram Singh · 66 Good"),
    ("Data Analyst", "Karan Patel · 80.7 Strong", "Priya Nair · 78 Good", "Aarav Sharma · 46.7 Partial"),
]
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        hdr = r == 0
        green = r > 0 and c == 1
        cell_set(tbl, r, c, val, size=14, bold=hdr or green,
                 color=EMERALD if green else WHITE,
                 fill=RGBColor(0x17, 0x20, 0x3F) if hdr else CARD)
bullets(s, [
    "Domain specialists top their own role — exactly what a recruiter would expect",
    "Fresh M.Tech ML grad ranks #3 on senior ML role (experience dimension at work)",
    "Every card opens an explainable report: score bars, ✓/✗ skills, gaps, recruiter note",
    "Agent Trace panel shows each agent's raw JSON — audit any decision live",
], t=4.5, size=16, gap=12)
footer(s, 7)

# ============================================================
# SLIDE 8 — INNOVATION & FAIRNESS
# ============================================================
s = slide()
title_bar(s, "07", "Innovation, Explainability & Fairness")
box(s, Inches(0.7), Inches(1.55), Inches(5.85), Inches(5.1), fill=CARD, border=INDIGO)
p = tb(s, Inches(0.95), Inches(1.7), Inches(5.4), Inches(0.5)).paragraphs[0]
run(p, "What makes us different", size=19, color=INDIGO, bold=True)
bullets(s, [
    "Hybrid LLM + deterministic design",
    "Dual-engine resilience (keyless demo)",
    "Adjacency-aware gap analysis",
    "Shared 130-skill taxonomy contract",
    "Recruiter-grade output, not a bare number",
    "One-click PS03-required demo",
], l=0.95, t=2.25, w=5.4, h=4.2, size=15, gap=12)
box(s, Inches(6.8), Inches(1.55), Inches(5.85), Inches(5.1), fill=CARD, border=EMERALD)
p = tb(s, Inches(7.05), Inches(1.7), Inches(5.4), Inches(0.5)).paragraphs[0]
run(p, "Trust & fairness by design", size=19, color=EMERALD, bold=True)
bullets(s, [
    "Published weights, plain-language notes",
    "Agent Trace = end-to-end audit",
    "No names / photos / ages in scoring",
    "Identical input ⇒ identical output",
    "Alias matching removes terminology bias",
    "Weights tunable per organisation",
], l=7.05, t=2.25, w=5.4, h=4.2, size=15, gap=12, bullet_color=EMERALD)
footer(s, 8)

# ============================================================
# SLIDE 9 — ROADMAP
# ============================================================
s = slide()
title_bar(s, "08", "Limitations & Roadmap")
bullets(s, [
    ("Today —", "heuristic offline extraction on very unstructured resumes; in-memory storage"),
    ("Next —", "SQLite/Postgres persistence, saved match history, recruiter accounts"),
    ("Next —", "semantic embeddings as a 6th scoring signal; layout-aware PDF parsing"),
    ("Next —", "multi-job comparison board; CSV/PDF export of ranked lists"),
    ("Next —", "scheduled fairness audits of weights & gap recommendations"),
], size=18, gap=16)
box(s, Inches(0.7), Inches(5.75), Inches(11.9), Inches(1.0), fill=RGBColor(0x17, 0x20, 0x3F), border=AMBER)
p = tb(s, Inches(0.95), Inches(5.9), Inches(11.4), Inches(0.7), anchor=MSO_ANCHOR.MIDDLE).paragraphs[0]
run(p, "Honest scope: we ship a complete, judge-proof vertical slice today — persistence & semantics are planned, not promised.", size=15, color=AMBER, bold=True)
footer(s, 9)

# ============================================================
# SLIDE 10 — TEAM / CLOSING
# ============================================================
s = slide()
p = tb(s, Inches(0.5), Inches(1.6), Inches(12.3), Inches(1.0)).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "Thank You", size=46, color=WHITE, bold=True)
p = tb(s, Inches(0.5), Inches(2.7), Inches(12.3), Inches(0.6)).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "[YOUR TEAM NAME]  —  questions welcome", size=20, color=INDIGO, bold=True)

box(s, Inches(2.6), Inches(3.6), Inches(8.1), Inches(2.3), fill=CARD, border=LINE)
tf = tb(s, Inches(2.85), Inches(3.75), Inches(7.6), Inches(2.0))
p = tf.paragraphs[0]
run(p, "AI & external resource disclosure: ", size=14, color=CYAN, bold=True)
run(p, "Google Gemini 2.5 Flash powers optional extraction & summaries; AI coding assistant used for scaffolding (all design decisions team-verified); demo data 100% synthetic; open datasets surveyed with license checks per PS03.", size=13, color=MUTED)
p2 = tf.add_paragraph()
p2.space_before = Pt(8)
run(p2, "Stack: React 18 · Vite · FastAPI · Gemini · python-docx/pypdf", size=13, color=MUTED)

p = tb(s, Inches(0.5), Inches(6.3), Inches(12.3), Inches(0.6)).paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "Live demo: load dataset → run pipeline → open any ranked candidate → inspect the Agent Trace", size=15, color=EMERALD, bold=True)
footer(s, 10)

prs.save("/home/user/docs/HireMind_AI_PitchDeck_PS03.pptx")
print("deck saved: 10 slides")
