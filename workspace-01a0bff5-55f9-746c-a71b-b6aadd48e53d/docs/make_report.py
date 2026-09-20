"""Generate the PS03 Documentation Round report (.docx)."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ---------- base style ----------
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

ACCENT = RGBColor(0x4F, 0x46, 0xE5)


def h1(t):
    p = doc.add_heading(t, level=1)
    return p


def h2(t):
    doc.add_heading(t, level=2)


def para(t, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    return p


def bullet(t):
    doc.add_paragraph(t, style="List Bullet")


def code(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = htext
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return t


# ============================================================
# COVER PAGE
# ============================================================
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("HireMind AI")
r.bold = True
r.font.size = Pt(40)
r.font.color.rgb = ACCENT

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Multi-Agent Resume Screening & Job Matching Platform")
r.font.size = Pt(16)

for _ in range(2):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Documentation Round Submission\nProblem Statement PS03")
r.font.size = Pt(13)
r.bold = True

doc.add_paragraph()
table(
    ["Field", "Details"],
    [
        ["Problem Statement", "PS03 — Multi-Agent Resume Screening & Job Matching Platform"],
        ["Technology Category", "NLP + LLM + Agentic AI"],
        ["Team Name", "[YOUR TEAM NAME]"],
        ["Team Members", "[Member 1] · [Member 2] · [Member 3]"],
        ["Hackathon", "[HACKATHON NAME]"],
        ["Submission Date", "11 September 2026"],
        ["Document Version", "1.0"],
    ],
)

doc.add_page_break()

# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================
h1("1. Executive Summary")
para(
    "HireMind AI is a multi-agent recruitment platform that automates the screening of resumes "
    "against job descriptions and produces a ranked, fully explainable list of candidate-role matches. "
    "The system is built around five cooperating software agents — a Resume Agent, a Job Agent, a "
    "Matching Agent, a Skill Gap Agent and a Recruiter Agent — mirroring the stages a human recruiting "
    "team would perform, but at machine speed and with full auditability."
)
para(
    "A key architectural decision distinguishes our solution from typical 'LLM-only' screeners: large "
    "language models (Google Gemini) are used only for *information extraction and natural-language "
    "summarisation*, while the actual ranking decision is produced by a *deterministic, weighted scoring "
    "engine*. This gives us the best of both worlds — rich semantic understanding of unstructured resumes, "
    "and reproducible, auditable, bias-controllable rankings. The platform runs end-to-end without any API "
    "key thanks to a deterministic offline NLP fallback layer, making it robust for live judging."
)
para(
    "The deliverable is a complete web application: a React 18 frontend served by a Python FastAPI "
    "backend. It satisfies the expected demonstration of PS03 — upload of 10 resumes and 3 job "
    "descriptions producing a ranked candidate list — and goes further with per-candidate explainability "
    "reports (score breakdown, matched/missing skills, upskilling recommendations, recruiter summary and a "
    "raw agent trace of every agent's input/output)."
)

# ============================================================
# 2. PROBLEM ANALYSIS
# ============================================================
h1("2. Problem Analysis")
para(
    "Recruiters spend an estimated 6-7 seconds on an initial resume scan, and high-volume roles routinely "
    "attract hundreds of applications per opening. Manual screening is slow, inconsistent between "
    "screeners, and prone to overlooking strong candidates whose resumes phrase skills differently from "
    "the job description ('React.js' vs 'React' vs 'ReactJS'). We decomposed the problem into five "
    "sub-problems:"
)
bullet("P1 — Heterogeneous input: resumes arrive as free text, PDFs or Word files with wildly different formats, section ordering and vocabularies.")
bullet("P2 — Requirement ambiguity: job descriptions mix hard requirements, preferences and marketing language; must-have vs nice-to-have skills must be separated programmatically.")
bullet("P3 — Compatibility measurement: skill overlap alone is insufficient — experience level, education and domain keywords all matter, with different importance per role.")
bullet("P4 — Actionable feedback: a bare score is not useful. Recruiters need to know *which* skills are missing and how close the candidate is; candidates benefit from upskilling guidance.")
bullet("P5 — Explainability and trust: screening affects livelihoods. Any automated ranking must justify itself per decision, must be reproducible, and must not silently depend on an opaque LLM opinion.")

para("These five sub-problems map one-to-one onto the five agents in our solution (Section 4).")

# ============================================================
# 3. RESEARCH
# ============================================================
h1("3. Research")
h2("3.1 Survey of existing approaches")
bullet("Keyword/ATS matching (legacy applicant-tracking systems): exact string matching on skills. Fast but brittle — fails on synonyms, abbreviations and formatting variance.")
bullet("Semantic embedding similarity: encode resume and JD into vectors and rank by cosine similarity. Handles synonyms well but produces a single opaque number; hard to explain 'why'.")
bullet("LLM-only screening: prompt an LLM to rate each candidate. Expressive but non-reproducible (scores drift between runs/models), expensive at scale, and difficult to audit for bias.")
bullet("Hybrid agentic systems (state of the art in 2025-26): decompose the workflow into specialised agents with tools, keeping a deterministic core for decisions. Our design follows this pattern.")

h2("3.2 Research conclusions that shaped our design")
bullet("Extraction is a language-understanding task → best solved by an LLM (Gemini), with rule-based fallback for reliability.")
bullet("Ranking is a *decision* task → must be deterministic, weighted and inspectable; weights are published constants.")
bullet("A curated skill taxonomy with aliases ('js'→javascript, 'k8s'→kubernetes, 'sklearn'→scikit-learn) plus fuzzy matching neutralises vocabulary mismatch — the single largest source of false negatives in keyword ATS systems.")

h2("3.3 Open datasets considered")
para(
    "Per the problem statement's source note, we surveyed public resume and job-description datasets "
    "available on open ML repositories (Kaggle's 'Resume Dataset' / 'Updated Resume Dataset' family and "
    "Hugging Face job-posting corpora). For the demonstration we deliberately ship a *synthetic* dataset "
    "of 10 resumes and 3 job descriptions (no real person's data), which also satisfies the 'no participant "
    "data collection' requirement. Should the team scale evaluation to real datasets before finals, each "
    "candidate dataset's license, version and access conditions will be verified and recorded before use, "
    "as the problem statement requires."
)

# ============================================================
# 4. PROPOSED SOLUTION
# ============================================================
h1("4. Proposed Solution")
para(
    "HireMind AI implements the five-agent pipeline shown below. Resumes and job descriptions are "
    "normalised by the extraction agents into a shared schema (canonical skills from a ~130-entry taxonomy, "
    "years of experience, education level). The Matching Agent computes a deterministic 0-100 fit score; "
    "the Skill Gap Agent converts misses into prioritised, actionable gaps; the Recruiter Agent writes the "
    "human summary and recommendation."
)
doc.add_picture("/home/user/docs/architecture.png", width=Inches(6.9))
last = doc.paragraphs[-1]
last.alignment = WD_ALIGN_PARAGRAPH.CENTER

h2("4.1 Agent responsibilities")
table(
    ["#", "Agent", "Input", "Output", "Engine"],
    [
        ["1", "Resume Agent", "Raw resume (txt/pdf/docx)", "Candidate profile: name, canonical skills, years of experience, education, highlights", "Gemini JSON-mode, offline heuristics fallback"],
        ["2", "Job Agent", "Raw job description", "Requirements: must-have skills, nice-to-have skills, min experience, education", "Gemini JSON-mode, offline heuristics fallback"],
        ["3", "Matching Agent", "Profile + requirements", "0-100 score, verdict, per-dimension breakdown, matched/missing skills", "Deterministic (always)"],
        ["4", "Skill Gap Agent", "Profile + match result", "Critical gaps, stretch gaps, adjacent strengths, upskilling plan", "Deterministic rule engine"],
        ["5", "Recruiter Agent", "Profile + requirements + match", "Recruiter-style summary + interview/hold/pass recommendation", "Gemini, template fallback"],
    ],
)

h2("4.2 Core design principles")
bullet("\"LLM extracts, never decides\" — language models parse and write; a transparent weighted engine ranks. Same input ⇒ same ranking, every time.")
bullet("Dual-mode operation — every LLM agent has a deterministic fallback, so the demo works offline, keyless, and never crashes mid-pitch.")
bullet("Single source of truth for skills — a shared taxonomy normalises both LLM output and heuristic scans; matching is engine-agnostic.")
bullet("Parallelism where latency matters — Recruiter summaries for all candidates are generated concurrently (thread pool of 5).")

# ============================================================
# 5. TECHNICAL FLOW
# ============================================================
h1("5. Technical Flow")
h2("5.1 End-to-end sequence")
bullet("Step 1 — Intake: recruiter uploads resumes (.txt/.pdf/.docx parsed server-side via pypdf / python-docx) or pastes text; job descriptions are pasted or loaded from the demo set.")
bullet("Step 2 — Extraction (Agents 1 & 2): each document is converted to the shared schema; if a Gemini API key is configured the LLM path runs (JSON-mode, schema-prompted), otherwise the offline NLP heuristics run (section scanning, taxonomy alias matching, regex experience/education extraction).")
bullet("Step 3 — Scoring (Agent 3): for every candidate×job pair, each dimension is scored 0-1 and combined with published weights; fuzzy matching (normalized containment + difflib ratio ≥ 0.85) links 'React.js' to 'react'.")
bullet("Step 4 — Gap analysis (Agent 4): missing must-haves become critical gaps with concrete upskilling recommendations; same-category candidate skills are surfaced as 'adjacent strengths' (e.g. missing Kubernetes but knows Docker).")
bullet("Step 5 — Summarisation (Agent 5): a recruiter-style paragraph with strengths, risks and a hire recommendation.")
bullet("Step 6 — Presentation: candidates are sorted into the ranked list; the UI exposes every intermediate artifact via an 'Agent Trace' panel.")

h2("5.2 Scoring model")
para("Score = Σ (weight × dimension score), renormalised when a dimension is not applicable:")
table(
    ["Dimension", "Weight", "Computation"],
    [
        ["Must-have skills", "45%", "fraction of required must-have skills matched (fuzzy)"],
        ["Nice-to-have skills", "15%", "fraction matched"],
        ["Experience", "20%", "1.0 if candidate ≥ required years; 0.6×(cand/req) if below; neutral 0.75 when unspecified"],
        ["Education", "10%", "degree-rank comparison (PhD > Master > Bachelor > Diploma); 0.8 when unspecified"],
        ["Keyword affinity", "10%", "JD keywords/titles found in resume text"],
    ],
)
para("Verdict bands: ≥80 Strong match · 60-79 Good match · 40-59 Partial match · <40 Weak match.")

h2("5.3 REST API surface")
table(
    ["Endpoint", "Purpose"],
    [
        ["POST /api/resumes/upload · POST /api/resumes", "Register resumes (files or pasted text) → Resume Agent runs"],
        ["POST /api/jobs", "Register a job description → Job Agent runs"],
        ["POST /api/demo/load", "Load bundled 10-resume / 3-JD demo set"],
        ["POST /api/match {job_id}", "Run Matching + Skill Gap + Recruiter agents over all candidates → ranked list"],
        ["POST /api/reextract", "Re-run extraction after adding a Gemini key"],
        ["GET /api/agents · /api/health · /api/settings", "Metadata, health, LLM configuration"],
    ],
)

# ============================================================
# 6. TECHNOLOGY STACK
# ============================================================
h1("6. Technology Stack")
table(
    ["Layer", "Technology", "Rationale"],
    [
        ["Frontend", "React 18 + Vite", "Component model ideal for agent cards, ranked lists, detail panels; Vite for fast HMR during the hackathon"],
        ["Backend", "Python 3.13 + FastAPI", "Async REST API, automatic OpenAPI docs, trivial JSON handling; served by Uvicorn"],
        ["LLM", "Google Gemini 2.5 Flash (REST, x-goog-api-key)", "JSON-mode structured extraction and summarisation; user-supplied key (never hard-coded)"],
        ["Offline NLP", "Regex + curated 130-skill taxonomy + difflib fuzzy matching", "Deterministic fallback; alias normalisation"],
        ["Document parsing", "pypdf, python-docx", "Accept real-world resume formats"],
        ["State", "In-memory store (swap-in SQLite/Postgres for production)", "Zero-setup demo; clean repository pattern for upgrade"],
        ["Concurrency", "ThreadPoolExecutor", "Parallel recruiter summaries"],
    ],
)

# ============================================================
# 7. IMPLEMENTATION
# ============================================================
h1("7. Implementation")
h2("7.1 Repository layout")
code(
    "backend/\n"
    "  main.py                FastAPI app: REST API + serves built React app\n"
    "  store.py               in-memory store\n"
    "  skills_taxonomy.py     canonical skills + alias map + categories\n"
    "  offline_extract.py     deterministic extraction heuristics\n"
    "  agents/\n"
    "    llm.py               Gemini REST client (JSON mode)\n"
    "    resume_agent.py      Agent 1\n"
    "    job_agent.py         Agent 2\n"
    "    matching_agent.py    Agent 3 (weights, fuzzy matching, verdicts)\n"
    "    skill_gap_agent.py   Agent 4 (gaps + upskilling hints)\n"
    "    recruiter_agent.py   Agent 5 (LLM / template summaries)\n"
    "  static/                built React frontend\n"
    "frontend/                React + Vite source (src/components/*)\n"
    "data/resumes/            10 synthetic demo resumes\n"
    "data/jobs/               3 synthetic demo job descriptions\n"
    "docs/                    this report + architecture diagram\n"
    "test_pipeline.py         end-to-end smoke test (load + match all jobs)"
)

h2("7.2 Demonstration data (expected demonstration of PS03)")
para(
    "The required demonstration — upload of 10 resumes and 3 job descriptions producing a ranked list — "
    "is available via one click (\"Load demo dataset\"). Resumes span ML, data science, frontend, "
    "full-stack, backend Java, DevOps, analytics, QA and mobile profiles; JDs cover Senior ML Engineer, "
    "Full-Stack Developer and Data Analyst. All demo data is synthetic."
)

h2("7.3 Verified results (offline mode, reproducible)")
para("The smoke test test_pipeline.py produced the following rankings:")
table(
    ["Job", "#1", "#2", "#3"],
    [
        ["Senior ML Engineer", "Aarav Sharma — 95.0 (Strong, 7/7 must-have)", "Priya Nair — 69.0 (Good)", "Divya Reddy — 59.2 (Partial, fresher penalty)"],
        ["Full-Stack Developer", "Sneha Kulkarni — 100.0 (Strong, 5/5)", "Rohan Mehta — 81.0 (Strong)", "Vikram Singh — 66.2 (Good)"],
        ["Data Analyst", "Karan Patel — 80.7 (Strong)", "Priya Nair — 78.0 (Good)", "Aarav Sharma — 46.7 (Partial)"],
    ],
)
para(
    "The rankings behave sensibly: domain specialists top their own role; a fresh M.Tech ML graduate ranks "
    "below experienced peers on the senior ML role (experience dimension); a frontend developer outranks a "
    "Java backend engineer for the full-stack role. Identical inputs reproduce identical outputs exactly."
)

h2("7.4 UI capabilities")
bullet("Tab 1 — Resumes: upload/paste; cards show extracted skills, experience, education and the extraction engine badge (Gemini vs offline).")
bullet("Tab 2 — Jobs: requirement chips (green = must-have, amber = nice-to-have), experience and education constraints.")
bullet("Tab 3 — Matches: job selector, animated 5-agent pipeline strip, ranked cards with per-dimension mini-bars; clicking opens the explainable report: recruiter summary, ✓/✗ skill coverage, weighted score bars, gap & upskilling panel, and the raw Agent Trace JSON.")
bullet("Settings: Gemini key + model (stored in browser), 'Re-extract with Gemini', full reset.")

# ============================================================
# 8. EXPLAINABILITY & FAIRNESS
# ============================================================
h1("8. Explainability & Fairness")
bullet("Every score decomposes into five published weighted components with plain-language notes (e.g. '3.5 yrs < 4 yrs required').")
bullet("The Agent Trace panel exposes each agent's engine and full JSON output — a match can be audited end-to-end.")
bullet("Scoring uses only skills, experience, education and JD keywords; no names, photos, genders, addresses or ages enter the scoring function.")
bullet("Deterministic core eliminates run-to-run variance; weights are constants in matching_agent.py and trivially adjustable per organisation.")
bullet("Fuzzy + taxonomy matching reduces systematic false-negatives that hurt candidates who use non-standard terminology.")

# ============================================================
# 9. INNOVATION
# ============================================================
h1("9. Innovation / What differentiates this solution")
bullet("Hybrid LLM+deterministic agent design ('LLM extracts, never decides') — explainability by construction, a pattern most LLM screeners lack.")
bullet("Dual-engine resilience: identical schema whether Gemini or offline heuristics ran; the demo cannot fail due to missing keys/network.")
bullet("Adjacency-aware gap analysis: instead of only listing missing skills, the Skill Gap Agent surfaces same-category strengths and a concrete learning path per gap.")
bullet("Shared skill taxonomy (130+ skills, aliases, categories) as the contract between all agents.")
bullet("Recruiter-grade output, not just a number: verdict bands, summary prose, interview/hold/pass recommendation.")
bullet("One-click, judge-proof demo dataset exactly matching the PS03 expected demonstration.")

# ============================================================
# 10. DEMO PLAN
# ============================================================
h1("10. Demo Plan (2-minute judging flow)")
bullet("0:00 — Load demo dataset; point out extraction badges and extracted skills on resume cards.")
bullet("0:30 — Tab 3, run the Matching Pipeline; narrate the five agent steps as the strip animates.")
bullet("0:50 — Show Senior ML Engineer ranking; open #1 (Aarav, 95) and a mid-rank candidate; explain score bars.")
bullet("1:20 — Open the Skill Gap panel: critical gaps + upskilling + adjacent strengths.")
bullet("1:40 — Open Agent Trace to demonstrate full auditability; switch job to show instant re-ranking.")

# ============================================================
# 11. LIMITATIONS & FUTURE WORK
# ============================================================
h1("11. Limitations & Future Work")
bullet("Offline extraction is heuristic; very unstructured resumes extract fewer skills than the LLM path. Mitigation: dual mode.")
bullet("Persistence is in-memory; roadmap: SQLite/Postgres, user accounts, saved match history.")
bullet("Roadmap: embedding-based semantic skill similarity as a sixth scoring signal; PDF-layout-aware parsing; side-by-side multi-job comparison board; CSV/PDF export of ranked lists for recruiters.")
bullet("Roadmap (fairness): periodic bias audits of weight choices and gap recommendations.")

# ============================================================
# 12. EXTERNAL / AI RESOURCES
# ============================================================
h1("12. Use of External / AI Resources (Disclosure)")
bullet("Google Gemini 2.5 Flash — used at runtime (optional, key supplied by user) for Resume/Job extraction and Recruiter summaries; accessed via official REST API.")
bullet("AI-assisted development — an agentic coding assistant (Arena.ai Agent Mode) was used to scaffold code, iterate on the UI and draft documentation text; all architectural decisions, scoring design and verification were reviewed and validated by the team.")
bullet("Open-source libraries — React, Vite, FastAPI, python-docx, pypdf, Pillow, difflib (see References).")
bullet("Demo data — fully synthetic, generated for this submission; no participant or personal data collected, per PS03.")
bullet("Public datasets — surveyed on Kaggle/Hugging Face for scale-up; none used without license verification (see 3.3).")

# ============================================================
# 13. REFERENCES
# ============================================================
h1("13. References")
bullet("[1] Google AI for Developers — Gemini API documentation. https://ai.google.dev")
bullet("[2] FastAPI documentation. https://fastapi.tiangolo.com")
bullet("[3] React documentation. https://react.dev")
bullet("[4] Vite documentation. https://vite.dev")
bullet("[5] Uvicorn — ASGI server. https://www.uvicorn.org")
bullet("[6] python-docx documentation. https://python-docx.readthedocs.io")
bullet("[7] pypdf documentation. https://pypdf.readthedocs.io")
bullet("[8] Python difflib — sequence matching. https://docs.python.org/3/library/difflib.html")
bullet("[9] Kaggle — Resume / job-description datasets (surveyed; verify license before use). https://www.kaggle.com/datasets")
bullet("[10] Hugging Face Datasets — job posting corpora (surveyed; verify license before use). https://huggingface.co/datasets")

doc.add_page_break()
h1("Appendix A — Run instructions")
code(
    "# backend (Python 3.10+)\n"
    "cd backend && pip install -r requirements.txt\n"
    "uvicorn main:app --host 0.0.0.0 --port 8000     # serves API + React build\n\n"
    "# frontend development (optional)\n"
    "cd frontend && npm install && npm run dev      # Vite, /api proxied to :8000\n\n"
    "# end-to-end smoke test\n"
    "python3 test_pipeline.py"
)
para(
    "Open http://localhost:8000 → 'Load demo dataset' → Tab 3 → 'Run Matching Pipeline'. "
    "Optionally add a Gemini key under Settings to switch extraction/summarisation to LLM mode."
)

doc.save("/home/user/docs/PS03_Documentation_HireMind_AI.docx")
print("report saved")
